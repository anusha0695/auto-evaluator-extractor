"""CLI: `python -m evaluation --doc <doc_id>`."""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path
from typing import Iterable

import yaml

from evaluation.lib.load_extraction import load_extraction
from evaluation.lib.load_ground_truth import load_ground_truth
from evaluation.lib.match_rows import match
from evaluation.lib.metrics import compute
from evaluation.lib.render_report import write_all, write_review_results_into

logger = logging.getLogger("evaluation")


def _repo_root() -> Path:
    """The repo root, derived from this file's location (evaluation/cli.py)."""
    return Path(__file__).resolve().parent.parent


def _default_paths(doc_id: str) -> tuple[Path, Path, Path]:
    """Convention paths (all relative to repo root)."""
    root = _repo_root()
    gt   = root / "ground_truth"   / f"{doc_id}.xlsx"
    ex   = root / "local_runs" / "artifacts"   / doc_id / "extraction_production.json"
    out  = root / "local_runs" / "evaluations" / doc_id
    return gt, ex, out


def _load_field_map() -> dict:
    p = _repo_root() / "evaluation" / "config" / "field_map.yaml"
    return yaml.safe_load(p.read_text(encoding="utf-8")) or {}


def _eval_one(
    doc_id: str,
    gt_path: Path,
    ex_path: Path,
    out_dir: Path,
    *,
    sheet_name: str | None = None,
    file_prefix: str = "",
    consolidated_wb=None,
    consolidated_sheet_name: str | None = None,
) -> int:
    """Run one eval. Optionally also append this doc's colored Review
    Results sheet to a shared `consolidated_wb` (the multi-sheet workbook
    flow uses this to build `<input>_evaluated.xlsx`)."""
    if not gt_path.is_file():
        logger.error("ground-truth file not found: %s", gt_path)
        return 2
    if not ex_path.is_file():
        logger.error("extraction file not found: %s", ex_path)
        return 2

    field_map = _load_field_map()
    gt_rows = load_ground_truth(gt_path, field_map, sheet_name=sheet_name)
    ex_rows = load_extraction(ex_path, field_map)

    logger.info("evaluation[%s]: gt_rows=%d  ex_rows=%d", doc_id, len(gt_rows), len(ex_rows))
    pairs = match(gt_rows, ex_rows)
    metrics = compute(pairs, field_map.get("columns") or [])

    # Peek the per-cell verdicts BEFORE write_all pops them — we may need
    # them again to render the consolidated workbook below.
    cell_verdicts_copy = list(metrics.get("_cell_verdicts", []) or [])

    write_all(out_dir, doc_id, gt_path, ex_path, metrics,
              field_map=field_map, file_prefix=file_prefix)
    logger.info("evaluation[%s]: wrote %s (prefix=%r)", doc_id, out_dir, file_prefix)

    # If the caller is also building a consolidated workbook, append this
    # doc's lab-layout review sheet to it. Same rendering function as the
    # per-doc mismatches.xlsx — same colors, same comments, same column
    # order — just under the doc filename as a tab inside the shared wb.
    if consolidated_wb is not None and cell_verdicts_copy:
        write_review_results_into(
            consolidated_wb, doc_id, cell_verdicts_copy,
            consolidated_sheet_name or doc_id,
            field_map=field_map,
        )

    return 0


def _discover_batch() -> Iterable[str]:
    """Doc IDs that have BOTH a GT.xlsx and an extraction_production.json."""
    root = _repo_root()
    gt_dir = root / "ground_truth"
    art_dir = root / "local_runs" / "artifacts"
    if not (gt_dir.is_dir() and art_dir.is_dir()):
        return []
    seen: set[str] = set()
    for gt_xlsx in gt_dir.glob("*.xlsx"):
        doc_id = gt_xlsx.stem
        if (art_dir / doc_id / "extraction_production.json").is_file():
            seen.add(doc_id)
    return sorted(seen)


# ─────────────────────────────────────────────────────────────────────────────
# Multi-sheet workbook mode — one workbook in, per-doc outputs into artifacts/
# ─────────────────────────────────────────────────────────────────────────────


def _doc_id_from_sheet_name(sheet_name: str) -> str:
    """Map an Excel sheet name to a doc_id.

    Excel sheet names can carry the source filename including its extension.
    Strip common doc extensions so a sheet named e.g. "full_report_Redacted.pdf"
    resolves to doc_id "full_report_Redacted" (which is what
    `local_runs/artifacts/<doc_id>/` is keyed by).
    """
    s = sheet_name.strip()
    for ext in (".pdf", ".PDF", ".xlsx", ".XLSX", ".docx", ".DOCX"):
        if s.endswith(ext):
            return s[: -len(ext)]
    return s


def _eval_workbook(workbook_path: Path) -> int:
    """For each sheet in `workbook_path`, treat the sheet name as a doc_id,
    find the matching `extraction_production.json` under
    `local_runs/artifacts/<doc_id>/`, run the eval, and write the 4 outputs
    (eval_metrics.json, eval_report.txt, eval_mismatches.xlsx, eval_summary.md)
    into that same artifacts dir.

    Sheets without a matching extraction are SKIPPED with a warning (not an
    error) so a workbook with index / summary sheets doesn't fail the run.

    ALSO builds a consolidated `<input_stem>_evaluated.xlsx` saved at
    `local_runs/evaluations/`. Each sheet in that consolidated workbook is
    ONE doc's color-coded lab-layout Review Results (the same view the SME
    sees in the per-doc `eval_mismatches.xlsx`), with the sheet name being
    the original input sheet name (the doc filename). The SME opens ONE
    file to scan every doc in the batch via sheet tabs.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("openpyxl is required: pip install openpyxl") from exc

    if not workbook_path.is_file():
        logger.error("workbook not found: %s", workbook_path)
        return 2

    wb = openpyxl.load_workbook(str(workbook_path), data_only=True, read_only=True)
    sheet_names = list(wb.sheetnames)
    logger.info("workbook[%s]: %d sheets — %s", workbook_path.name,
                len(sheet_names), sheet_names)

    root = _repo_root()
    art_dir = root / "local_runs" / "artifacts"

    # Consolidated workbook — accumulated across all docs in this batch.
    # We'll remove the default "Sheet" at the end (or before saving) so the
    # output contains only the per-doc review sheets.
    consolidated_wb = openpyxl.Workbook()

    rc = 0
    n_ran = 0
    n_skipped = 0
    for sheet in sheet_names:
        doc_id = _doc_id_from_sheet_name(sheet)
        ex_path = art_dir / doc_id / "extraction_production.json"
        if not ex_path.is_file():
            logger.warning(
                "sheet %r → doc_id %r: no extraction at %s, skipping",
                sheet, doc_id, ex_path)
            n_skipped += 1
            continue
        out_dir = art_dir / doc_id
        # Pass the original sheet name so load_ground_truth picks the right tab,
        # AND so the consolidated workbook's tab matches the SME's input name.
        single_rc = _eval_one(
            doc_id, workbook_path, ex_path, out_dir,
            sheet_name=sheet, file_prefix="eval_",
            consolidated_wb=consolidated_wb,
            consolidated_sheet_name=sheet,   # keep the user's original sheet name
        )
        rc |= single_rc
        n_ran += 1

    # Drop the default empty "Sheet" openpyxl created, but only if we wrote
    # at least one real sheet (else save would fail on a zero-sheet wb).
    default_sheet = "Sheet"
    if default_sheet in consolidated_wb.sheetnames and len(consolidated_wb.sheetnames) > 1:
        del consolidated_wb[default_sheet]

    # Save the consolidated workbook at local_runs/evaluations/<input>_evaluated.xlsx
    evals_dir = root / "local_runs" / "evaluations"
    evals_dir.mkdir(parents=True, exist_ok=True)
    out_xlsx = evals_dir / f"{workbook_path.stem}_evaluated.xlsx"
    if n_ran == 0:
        logger.warning(
            "consolidated workbook NOT written: no doc sheets matched any "
            "extraction_production.json (intended path: %s)", out_xlsx)
    else:
        consolidated_wb.save(str(out_xlsx))
        logger.info("consolidated workbook: wrote %s (%d sheet(s))",
                    out_xlsx, len(consolidated_wb.sheetnames))

    logger.info("workbook done: ran=%d  skipped=%d  rc=%d", n_ran, n_skipped, rc)
    return rc


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        prog="python -m evaluation",
        description=(
            "Evaluate extraction_production.json against SME ground-truth XLSX. "
            "Three modes: (1) --doc / --batch reads a single-sheet GT per doc "
            "from ground_truth/<doc>.xlsx and writes to local_runs/evaluations/. "
            "(2) --workbook reads a multi-sheet GT workbook (sheet name = doc "
            "filename) and writes per-doc outputs into local_runs/artifacts/<doc>/ "
            "with an 'eval_' filename prefix so they coexist with pipeline "
            "artifacts (extraction_production.json, link_metrics.json, etc)."
        ),
    )
    ap.add_argument("--doc", help="doc_id (e.g. full_report_Redacted). Required unless --batch or --workbook.")
    ap.add_argument("--gt", type=Path, default=None,
                    help="Override path to GT xlsx (default: ground_truth/<doc>.xlsx)")
    ap.add_argument("--extraction", type=Path, default=None,
                    help="Override path to extraction_production.json")
    ap.add_argument("--out", type=Path, default=None,
                    help="Override output dir (default: local_runs/evaluations/<doc>/)")
    ap.add_argument("--batch", action="store_true",
                    help="Eval every doc with both GT.xlsx + extraction_production.json")
    ap.add_argument("--workbook", type=Path, default=None,
                    help=("Multi-sheet GT workbook. Each sheet name = doc filename "
                          "(extension stripped). Outputs land in "
                          "local_runs/artifacts/<doc>/ with 'eval_' prefix."))
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)s [%(name)s] %(message)s",
    )

    if args.workbook:
        return _eval_workbook(args.workbook)

    if args.batch:
        rc = 0
        for doc_id in _discover_batch():
            gt, ex, out = _default_paths(doc_id)
            rc |= _eval_one(doc_id, gt, ex, out)
        return rc

    if not args.doc:
        ap.error("--doc is required (or use --batch / --workbook)")

    gt, ex, out = _default_paths(args.doc)
    if args.gt:         gt = args.gt
    if args.extraction: ex = args.extraction
    if args.out:        out = args.out
    return _eval_one(args.doc, gt, ex, out)


if __name__ == "__main__":
    sys.exit(main())
