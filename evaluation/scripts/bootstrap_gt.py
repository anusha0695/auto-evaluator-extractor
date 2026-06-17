"""Bootstrap a ground-truth XLSX from a Python dict literal.

Usage:
    python evaluation/scripts/bootstrap_gt.py --doc full_report_Redacted

Writes:
    ground_truth/<doc>.xlsx

The default content for `full_report_Redacted` is encoded inline below — values
verbatim from the source PDF, with the page-6/7 appendix values winning over
page-1 summary values per the supersession rule we agreed on. Edit the
`GT_CONTENT` constant to update or add new documents.

The output workbook uses the 20-column lab layout (matches the reference
screenshots). Sheet name: "Review Results". Header row in row 1; data rows
starting row 2. One sheet, variants stacked above biomarkers — the eval
module does match-against-both row pairing, so order doesn't matter.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


def _repo_root() -> Path:
    return Path(__file__).resolve().parent.parent.parent


# 20-column lab layout (matches evaluation/config/field_map.yaml and the
# reference screenshots). Duplicate "Amino Acid Change" preserved per lab convention.
COLUMNS = [
    "Biomarker", "Method", "Result", "Interpretation", "Reference Range",
    "Variant Allele Freq", "DNA Change Type", "Amino Acid Change",
    "Genomic DNA Change", "Genomic Ref Seq", "Coding DNA Change",
    "Transcript Ref", "Amino Acid Change", "AA Ref Seq",
    "Clinical Significance", "Genomic Source", "Ref Assembly",
    "Chromosome", "Genomic Position", "Exon",
]


# Per-doc GT content. Each row is a dict keyed by column header. Missing keys
# render as blank cells. For duplicate columns ("Amino Acid Change" appears
# twice) the SAME value is written to both cells.
GT_CONTENT: dict[str, list[dict]] = {

    "full_report_Redacted": [

        # ---- VARIANTS (winner = page-6/7 appendix per supersession) ----
        {
            "Biomarker": "KRAS", "Method": "NGS", "Result": "Pathogenic Variant",
            "Variant Allele Freq": "10",
            "DNA Change Type": "Substitution",
            "Amino Acid Change": "p.G12D",
            "Coding DNA Change": "c.35G>A",
            "Transcript Ref": "NM_004985.4",
            "Clinical Significance": "Pathogenic Variant",
            "Genomic Source": "Somatic",
            "Chromosome": "chr12",
            "Genomic Position": "25245350",
            "Exon": "2",
        },
        {
            "Biomarker": "TP53", "Method": "NGS", "Result": "Pathogenic Variant",
            "Variant Allele Freq": "7",
            "DNA Change Type": "Substitution",
            "Amino Acid Change": "p.C242F",
            "Coding DNA Change": "c.725G>T",
            "Clinical Significance": "Pathogenic Variant",
            "Genomic Source": "Somatic",
            "Chromosome": "chr17",
            "Genomic Position": "7674238",
            "Exon": "7",
        },
        {
            "Biomarker": "STK11", "Method": "NGS",
            "Result": "Variant of Uncertain Significance",
            "Variant Allele Freq": "6",
            "DNA Change Type": "Substitution",
            "Amino Acid Change": "p.L190R",
            "Coding DNA Change": "c.569T>G",
            "Transcript Ref": "NM_000455.4",
            "Clinical Significance": "Variant of Uncertain Significance",
            "Genomic Source": "Somatic",
            "Chromosome": "chr19",
            "Genomic Position": "1220477",
            "Exon": "4",
        },
        {
            "Biomarker": "ASXL1", "Method": "NGS", "Result": "Pathogenic Variant",
            "Variant Allele Freq": "0.3%",
            "DNA Change Type": "Deletion",
            "Amino Acid Change": "P808fs",
            "Coding DNA Change": "c.2423delC",
            "Clinical Significance": "Pathogenic Variant",
            "Genomic Source": "Somatic",
        },
        {
            "Biomarker": "EGFR", "Method": "NGS", "Result": "Likely Pathogenic Variant",
            "Variant Allele Freq": "0.2%",
            "DNA Change Type": "Substitution",
            "Amino Acid Change": "P596L",
            "Coding DNA Change": "c.1787C>T",
            "Clinical Significance": "Likely Pathogenic Variant",
            "Genomic Source": "Somatic",
        },
        {
            "Biomarker": "LZTR1", "Method": "NGS", "Result": "Pathogenic Variant",
            "Variant Allele Freq": "0.1%",
            "DNA Change Type": "Substitution",
            "Amino Acid Change": "C342R",
            "Coding DNA Change": "c.1024T>C",
            "Clinical Significance": "Pathogenic Variant",
            "Genomic Source": "Somatic",
        },

        # ---- BIOMARKERS (other_molecular_biomarker_umbrella) ----
        {
            "Biomarker": "BLOOD TMB (mut/Mb)", "Method": "NGS", "Result": "5",
        },
        {
            "Biomarker": "MICROSATELLITE INSTABILITY", "Method": "NGS",
            "Result": "Not Detected",
        },
        {
            "Biomarker": "TUMOR FRACTION", "Method": "NGS", "Result": "9.8%",
        },
        {
            "Biomarker": "HLA-A", "Method": "Seq",
            "Result": "A*30:02, A*02:01",
        },
        {
            "Biomarker": "HLA-B", "Method": "Seq",
            "Result": "B*44:02, B*18:01",
        },
        {
            "Biomarker": "HLA-C", "Method": "Seq",
            "Result": "C*05:01, C*05:01",
        },
    ],
}


def write_gt(doc_id: str, out_path: Path) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError as exc:
        raise SystemExit("openpyxl required: pip install openpyxl") from exc

    rows = GT_CONTENT.get(doc_id)
    if rows is None:
        raise SystemExit(
            f"No GT_CONTENT entry for doc_id={doc_id!r}. Add a row list to "
            f"GT_CONTENT in evaluation/scripts/bootstrap_gt.py."
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Review Results"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for col_idx, header in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align

    # Sensible column widths (the lab review uses similar widths).
    widths = {
        "A": 24, "B": 8, "C": 24, "D": 22, "E": 18, "F": 14, "G": 14, "H": 16,
        "I": 18, "J": 18, "K": 18, "L": 14, "M": 16, "N": 14, "O": 28,
        "P": 16, "Q": 16, "R": 12, "S": 16, "T": 8,
    }
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w
    ws.freeze_panes = "A2"

    for r_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(COLUMNS, start=1):
            ws.cell(row=r_idx, column=col_idx, value=row.get(header))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))


def main() -> int:
    ap = argparse.ArgumentParser(description="Generate a ground-truth XLSX from inline content.")
    ap.add_argument("--doc", required=True,
                    help="doc_id (must match a key in GT_CONTENT, e.g. full_report_Redacted)")
    ap.add_argument("--out", type=Path, default=None,
                    help="Override output path (default: ground_truth/<doc>.xlsx)")
    args = ap.parse_args()

    out = args.out or (_repo_root() / "ground_truth" / f"{args.doc}.xlsx")
    write_gt(args.doc, out)
    print(f"Wrote: {out}")
    print(f"Rows:  {len(GT_CONTENT[args.doc])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
